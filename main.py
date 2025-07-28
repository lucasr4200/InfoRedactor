from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from gliner import GLiNER
from typing import List
import fitz  # PyMuPDF
import pdfplumber
from docx import Document
import io
from io import BytesIO

#Load model
print("Loading model...")
model = GLiNER.from_pretrained("urchade/gliner_small-v2.1")
print("Model loaded")

app = FastAPI()

#Serve static files
app.mount("/static", StaticFiles(directory="static"), name="static")

#Define labels
labels = ["email address", "person", "organization", "location", "insurance", "address", "postal code",
          "employee id", "passport number", "phone number", "credit card number", "credit card cvv",
          "credit card expiry", "bank account number", "driver's license number", "social insurance number",
          "date of birth", "password"]


#Define Pydantic model for input data
class TextInput(BaseModel):
    text: str


def redact_text(text: str, entities: List[dict]) -> str:
    # Sort entities by their start position in reverse order
    sorted_entities = sorted(entities, key=lambda e: e['start'], reverse=True)

    # Redact text
    redacted_text = text
    for entity in sorted_entities:
        start = entity['start']
        end = entity['end']
        label = f"<{entity['label']}>"
        redacted_text = redacted_text[:start] + label + redacted_text[end:]

    return redacted_text


#For index.html
@app.get("/", response_class=HTMLResponse)
async def get_index():
    with open("static/index.html", "r") as f:
        return HTMLResponse(content=f.read())


#Route to handle redaction
@app.post("/redact/")
async def redact_text_endpoint(input_data: TextInput):
    #Predict entities
    predicted_entities = model.predict_entities(input_data.text, labels, threshold=0.15)
    #Redact input text based on predicted entities
    redacted_text = redact_text(input_data.text, predicted_entities)
    return {"redacted_text": redacted_text}


#Function to process and redact PDF files
def process_pdf(file_stream: BytesIO) -> BytesIO:
    try:
        output = BytesIO()
        #Open PDF file
        with fitz.open(stream=file_stream, filetype="pdf") as pdf:
            #Create new PDF for redacted text
            new_pdf = fitz.open()

            #Process each page
            for page_num in range(len(pdf)):
                page = pdf.load_page(page_num)
                text = page.get_text("text")  # Extract text

                #Predict and redact text
                predicted_entities = model.predict_entities(text, labels, threshold=0.15)
                redacted_text = redact_text(text, predicted_entities)

                #Create a new page in new PDF
                new_page = new_pdf.new_page(width=page.rect.width, height=page.rect.height)

                #Add redacted text to new page
                new_page.insert_text((72, 72), redacted_text, fontsize=11, color=(0, 0, 0))

            #Save new PDF to output
            new_pdf.save(output)

        output.seek(0)
        return output

    except Exception as e:
        print(f"Error processing PDF: {e}")
        raise HTTPException(status_code=500, detail="An error occurred while processing the PDF file.")


#Process and redact word files
def process_word(file_stream: io.BytesIO) -> io.BytesIO:
    output = io.BytesIO()
    doc = Document(file_stream)
    new_doc = Document()

    for para in doc.paragraphs:
        text = para.text
        predicted_entities = model.predict_entities(text, labels, threshold=0.15)
        redacted_text = redact_text(text, predicted_entities)
        new_doc.add_paragraph(redacted_text)

    new_doc.save(output)
    output.seek(0)
    return output

#Handle PDF file upload and redaction
@app.post("/upload_pdf/")
async def upload_pdf(file: UploadFile = File(...)):
    if not file.filename.endswith('.pdf'):
        raise HTTPException(status_code=400, detail="Invalid file type. Only PDF files are supported.")
    file_stream = io.BytesIO(await file.read())
    redacted_pdf = process_pdf(file_stream)
    return StreamingResponse(redacted_pdf, media_type="application/pdf",
                             headers={"Content-Disposition": "attachment; filename=redacted.pdf"})


#Handle word file upload and redaction
@app.post("/upload_word/")
async def upload_word(file: UploadFile = File(...)):
    if not file.filename.endswith('.docx'):
        raise HTTPException(status_code=400, detail="Invalid file type. Only DOCX files are supported.")
    file_stream = io.BytesIO(await file.read())
    redacted_word = process_word(file_stream)
    return StreamingResponse(redacted_word,
                             media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                             headers={"Content-Disposition": "attachment; filename=redacted.docx"})